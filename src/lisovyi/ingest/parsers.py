"""Текст документа з його байтів.

Функції приймають байти, а не шлях: той самий код оброблятиме файли, які надсилають звірі.
"""

import hashlib
import io
import json
import logging
from dataclasses import dataclass, field
from pathlib import Path, PurePath

import docx
import pypdf
from docx.table import Table, _Cell

from lisovyi.domain import TextSource
from lisovyi.ingest import ocr
from lisovyi.ingest.ocr import ReviewItem

log = logging.getLogger(__name__)


class UnsupportedDocument(ValueError):
    """Файл, з якого не вдається отримати текст: невідомий формат, кодування чи пошкоджений."""


class NoTextFound(UnsupportedDocument):
    """Файл прочитано, але тексту в ньому немає: навіть OCR нічого не знайшов."""


@dataclass(frozen=True)
class Extracted:
    text: str
    source: TextSource
    # Тільки для source=ocr: місця, які варто перепитати у звіра.
    review_items: list[ReviewItem] = field(default_factory=list)


@dataclass(frozen=True)
class Transcript:
    text: str
    verified: bool


class TranscriptStore:
    """Вивірені вручну транскрипти сканів, прив'язані до sha256 PDF-файлу."""

    def __init__(self, by_sha256: dict[str, Transcript]):
        self._by_sha256 = by_sha256

    @classmethod
    def load(cls, directory: Path) -> "TranscriptStore":
        manifest = json.loads((directory / "manifest.json").read_text(encoding="utf-8"))
        by_sha256 = {}
        for pdf_name, entry in manifest.items():
            path = directory / entry["transcript"]
            if not path.is_file():
                raise FileNotFoundError(f"transcript for {pdf_name!r} is missing: {path}")
            by_sha256[entry["sha256"]] = Transcript(
                path.read_text(encoding="utf-8"), bool(entry["verified"])
            )
        return cls(by_sha256)

    def get(self, sha256: str) -> Transcript | None:
        return self._by_sha256.get(sha256)


def extract_text(data: bytes, filename: str, transcripts: TranscriptStore | None) -> Extracted:
    suffix = PurePath(filename).suffix.lower()
    if suffix == ".txt":
        return Extracted(_decode_txt(data, filename), TextSource.TEXT)
    if suffix == ".docx":
        return Extracted(_docx_text(data, filename), TextSource.TEXT)
    if suffix == ".pdf":
        return _pdf_text(data, filename, transcripts)
    raise UnsupportedDocument(f"{filename}: unsupported file type {suffix or '(none)'!r}")


def _decode_txt(data: bytes, filename: str) -> str:
    # cp1251 декодує майже будь-які байти; нульові байти означають UTF-16 або бінарний
    # файл, і без цієї перевірки вони стали б мовчазними «кракозябрами».
    if b"\x00" in data:
        raise UnsupportedDocument(f"{filename}: text contains NUL bytes (UTF-16 or binary)")
    # utf-8-sig — це UTF-8, що прибирає BOM, який дописує Блокнот.
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise UnsupportedDocument(f"{filename}: text is neither UTF-8 nor cp1251")


def _docx_text(data: bytes, filename: str) -> str:
    # python-docx і lxml падають на зіпсованих файлах дуже різними винятками (BadZipFile,
    # KeyError, XMLSyntaxError, ValueError…): будь-який із них означає «не можу прочитати».
    try:
        document = docx.Document(io.BytesIO(data))
        return "\n".join(_blocks_text(document.iter_inner_content()))
    except Exception as exc:
        raise UnsupportedDocument(f"{filename}: not a readable DOCX ({exc!r})") from exc


def _blocks_text(blocks) -> list[str]:
    """Абзаци й таблиці в порядку документа."""
    lines = []
    for block in blocks:
        if isinstance(block, Table):
            lines.extend(_table_lines(block))
        else:
            lines.append(block.text)
    return lines


def _table_lines(table: Table) -> list[str]:
    lines = []
    for row in table.rows:
        cells: list[_Cell] = []
        for cell in row.cells:
            # Об'єднана клітинка повторюється в row.cells для кожної колонки, яку займає.
            if not cells or cell._tc is not cells[-1]._tc:
                cells.append(cell)
        lines.append(" | ".join(_cell_text(cell) for cell in cells))
    return lines


def _cell_text(cell: _Cell) -> str:
    # cell.text бере лише абзаци; таблиця, вкладена в клітинку, з нього випадає.
    return "\n".join(_blocks_text(cell.iter_inner_content()))


def _pdf_text(data: bytes, filename: str, transcripts: TranscriptStore | None) -> Extracted:
    # pypdf на битих файлах кидає не лише PyPdfError, а й ValueError, KeyError, TypeError…
    try:
        reader = pypdf.PdfReader(io.BytesIO(data))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as exc:
        raise UnsupportedDocument(f"{filename}: not a readable PDF ({exc!r})") from exc
    if text.strip():
        return Extracted(text, TextSource.TEXT)

    # Скан без текстового шару.
    if transcripts is not None:
        transcript = transcripts.get(hashlib.sha256(data).hexdigest())
        if transcript is not None:
            if not transcript.verified:
                log.warning("%s: using a transcript that is not verified", filename)
            return Extracted(transcript.text, TextSource.TRANSCRIPT)

    text = ocr.ocr_pdf(data)
    if not text.strip():
        raise NoTextFound(f"{filename}: OCR found no text")
    return Extracted(text, TextSource.OCR, ocr.find_review_items(text))
