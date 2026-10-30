import io
import json
import logging
import shutil
import zipfile

import docx
import pypdf
import pytest

from lisovyi.config import PROJECT_ROOT
from lisovyi.domain import TextSource
from lisovyi.ingest import ocr
from lisovyi.ingest.parsers import (
    NoTextFound,
    TranscriptStore,
    UnsupportedDocument,
    extract_text,
)

ARCHIVE = PROJECT_ROOT / "data" / "archive"
TRANSCRIPTS = PROJECT_ROOT / "data" / "transcripts"
SCAN = "Як утеплити рукавичку біля річки.pdf"


def test_txt_utf8():
    result = extract_text("Правила толоки".encode(), "a.txt", None)
    assert (result.text, result.source) == ("Правила толоки", TextSource.TEXT)


def test_txt_falls_back_to_cp1251():
    data = "Ця пам'ятка для всіх мешканців лісу".encode("cp1251")
    with pytest.raises(UnicodeDecodeError):
        data.decode("utf-8")
    assert extract_text(data, "a.txt", None).text == "Ця пам'ятка для всіх мешканців лісу"


def test_txt_in_neither_encoding_is_unsupported():
    # 0x98 не має символу в cp1251 і не є коректним UTF-8.
    with pytest.raises(UnsupportedDocument, match="neither UTF-8 nor cp1251"):
        extract_text(b"\x98", "a.txt", None)


def test_txt_utf8_bom_is_dropped():
    assert extract_text("Правила толоки".encode("utf-8-sig"), "a.txt", None).text == (
        "Правила толоки"
    )


def test_txt_in_utf16_is_unsupported():
    # Без перевірки cp1251 «успішно» декодував би ці байти в сміття з NUL.
    with pytest.raises(UnsupportedDocument, match="NUL bytes"):
        extract_text("Правила толоки".encode("utf-16"), "a.txt", None)


def _docx_bytes(build) -> bytes:
    document = docx.Document()
    build(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_docx_includes_tables_in_document_order():
    def build(document):
        document.add_paragraph("Перед таблицею")
        table = document.add_table(rows=2, cols=2)
        for row, values in zip(
            table.rows, [("Нора", "3 шишки"), ("Дупло", "2 жолуді")], strict=True
        ):
            for cell, value in zip(row.cells, values, strict=True):
                cell.text = value
        document.add_paragraph("Після таблиці")

    result = extract_text(_docx_bytes(build), "a.docx", None)
    assert result.source == TextSource.TEXT
    assert result.text.split("\n") == [
        "Перед таблицею",
        "Нора | 3 шишки",
        "Дупло | 2 жолуді",
        "Після таблиці",
    ]


def test_docx_merged_cell_is_not_repeated():
    def build(document):
        table = document.add_table(rows=1, cols=3)
        merged = table.cell(0, 0).merge(table.cell(0, 1))
        merged.text = "Внесок"
        table.cell(0, 2).text = "3 шишки"

    assert extract_text(_docx_bytes(build), "a.docx", None).text == "Внесок | 3 шишки"


def test_docx_nested_table_text_is_collected():
    def build(document):
        outer = document.add_table(rows=1, cols=2)
        outer.cell(0, 0).text = "Нора"
        cell = outer.cell(0, 1)
        cell.text = "Внески:"
        inner = cell.add_table(rows=1, cols=2)
        inner.cell(0, 0).text = "дорослий"
        inner.cell(0, 1).text = "3 шишки"

    text = extract_text(_docx_bytes(build), "a.docx", None).text
    # Word вимагає абзац після вкладеної таблиці, тож клітинка закінчується порожнім рядком.
    assert text.split("\n") == ["Нора | Внески:", "дорослий | 3 шишки", ""]


@pytest.mark.parametrize("error", [RuntimeError, ValueError, AttributeError, KeyError])
def test_any_docx_library_error_is_unsupported(error, monkeypatch):
    def broken(stream):
        raise error("boom")

    monkeypatch.setattr(docx, "Document", broken)
    with pytest.raises(UnsupportedDocument, match="not a readable DOCX"):
        extract_text(b"PK", "a.docx", None)


@pytest.mark.parametrize("error", [RuntimeError, ValueError, KeyError, TypeError])
def test_any_pdf_library_error_is_unsupported(error, monkeypatch):
    def broken(stream):
        raise error("boom")

    monkeypatch.setattr(pypdf, "PdfReader", broken)
    with pytest.raises(UnsupportedDocument, match="not a readable PDF"):
        extract_text(b"%PDF-1.4", "a.pdf", None)


@pytest.mark.parametrize("filename", ["Порада.doc", "Порада.rtf", "Порада"])
def test_other_extensions_are_unsupported(filename):
    with pytest.raises(UnsupportedDocument, match="unsupported file type"):
        extract_text(b"whatever", filename, None)


@pytest.mark.parametrize("filename", ["a.docx", "a.pdf"])
def test_corrupt_files_are_unsupported(filename):
    with pytest.raises(UnsupportedDocument, match="not a readable"):
        extract_text(b"not really a document", filename, None)


def test_docx_with_malformed_xml_is_unsupported():
    source = zipfile.ZipFile(io.BytesIO(_docx_bytes(lambda document: None)))
    broken = io.BytesIO()
    with zipfile.ZipFile(broken, "w") as target:
        for entry in source.infolist():
            data = source.read(entry)
            target.writestr(
                entry, b"<w:document" if entry.filename == "word/document.xml" else data
            )

    with pytest.raises(UnsupportedDocument, match="not a readable DOCX"):
        extract_text(broken.getvalue(), "a.docx", None)


def test_pdf_with_text_layer():
    result = extract_text((ARCHIVE / "Карта лісу.pdf").read_bytes(), "Карта лісу.pdf", None)
    assert result.source == TextSource.TEXT
    assert "Легенда карти" in result.text
    assert result.review_items == []


def test_scan_uses_transcript_matched_by_sha256():
    result = extract_text((ARCHIVE / SCAN).read_bytes(), SCAN, TranscriptStore.load(TRANSCRIPTS))
    assert result.source == TextSource.TRANSCRIPT
    expected = (TRANSCRIPTS / "Як утеплити рукавичку біля річки.txt").read_text(encoding="utf-8")
    assert result.text == expected


def _copy_transcripts(tmp_path, change) -> TranscriptStore:
    directory = tmp_path / "transcripts"
    shutil.copytree(TRANSCRIPTS, directory)
    manifest_path = directory / "manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    change(manifest[SCAN])
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
    return TranscriptStore.load(directory)


@pytest.fixture
def fake_ocr(monkeypatch):
    calls = []

    def ocr_pdf(data: bytes) -> str:
        calls.append(data)
        return "Внесок становить З шишки."

    monkeypatch.setattr(ocr, "ocr_pdf", ocr_pdf)
    return calls


def test_transcript_is_ignored_when_sha256_differs(tmp_path, fake_ocr):
    store = _copy_transcripts(tmp_path, lambda entry: entry.update(sha256="0" * 64))
    data = (ARCHIVE / SCAN).read_bytes()

    result = extract_text(data, SCAN, store)

    assert fake_ocr == [data]
    assert result.source == TextSource.OCR
    assert result.text == "Внесок становить З шишки."
    assert [(i.token, i.reason) for i in result.review_items] == [("З", "digit_like_letter")]


def test_scan_without_transcript_store_goes_to_ocr(fake_ocr):
    result = extract_text((ARCHIVE / SCAN).read_bytes(), SCAN, None)
    assert result.source == TextSource.OCR
    assert len(fake_ocr) == 1


def test_unverified_transcript_is_used_with_warning(tmp_path, caplog, fake_ocr):
    store = _copy_transcripts(tmp_path, lambda entry: entry.update(verified=False))

    with caplog.at_level(logging.WARNING):
        result = extract_text((ARCHIVE / SCAN).read_bytes(), SCAN, store)

    assert result.source == TextSource.TRANSCRIPT
    assert fake_ocr == []
    assert [r.levelno for r in caplog.records] == [logging.WARNING]
    assert SCAN in caplog.text


def test_ocr_without_any_text_is_unsupported(monkeypatch):
    monkeypatch.setattr(ocr, "ocr_pdf", lambda data: " \n\n ")
    with pytest.raises(NoTextFound, match="OCR found no text"):
        extract_text((ARCHIVE / SCAN).read_bytes(), SCAN, None)


def test_transcript_store_fails_on_missing_transcript_file(tmp_path):
    with pytest.raises(FileNotFoundError, match="transcript for"):
        _copy_transcripts(tmp_path, lambda entry: entry.update(transcript="немає.txt"))
