"""Тестові поради звірів для брифу 04. Результати закомічені поруч, тож тестам шрифт не потрібен.

Скан рендериться кириличним шрифтом, якого в образі python:3.12-slim немає. У WSL:

    sudo apt install fonts-dejavu-core
    docker compose run --rm -v "$PWD:/app" -v /usr/share/fonts/truetype/dejavu:/fonts:ro \\
        bot python tests/fixtures/uploads/make_fixtures.py /fonts/DejaVuSans.ttf
"""

import argparse
import io
import random
import re
from pathlib import Path

import docx
import pypdf
from PIL import Image, ImageDraw, ImageFont

from lisovyi.config import PROJECT_ROOT
from lisovyi.ingest.parsers import TranscriptStore, extract_text

HERE = Path(__file__).parent
ARCHIVE = PROJECT_ROOT / "data" / "archive"
GUIDE = ARCHIVE / "Грибний путівник хвойного лісу.pdf"
MAP = ARCHIVE / "Карта лісу.pdf"
ON_EDGE = "На узліссі — ні хвойний, ні листяний ліс, і до річки далеко."

VERBA_ADVICE = (
    "Мешканцям рукавички «Під вербою». Навесні латаття біля берега цвіте, і туди прилітають "
    "бабки. Малим їх не ловити."
)
GUIDE_INTRO = "Цей путівник написано для звірів, що живуть у хвойному лісі."
GUIDE_NEW_SENTENCE = "Після першого снігу грибів у хвойному лісі вже не шукайте."
NO_PLACE = "Корисна порада. Ягоди збирайте зранку, поки роса."
SCAN_LINES = ["До дуба — 300 кроків.", "Внесок — 3 шишки."]
SEED = 4


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("font", type=Path, help="TTF з кирилицею, напр. DejaVuSans.ttf")
    font = parser.parse_args().font

    (HERE / "порада.txt").write_text(VERBA_ADVICE + "\n", encoding="utf-8")
    (HERE / "без_місця.txt").write_text(NO_PLACE + "\n", encoding="utf-8")
    (HERE / "гриби_нова_версія.docx").write_bytes(_mushroom_version())
    (HERE / "скан_з_числами.pdf").write_bytes(_scan(font))
    (HERE / "Карта лісу.txt").write_text(_map(), encoding="utf-8")
    (HERE / "задовга.txt").write_text(_too_long(), encoding="utf-8")
    (HERE / "порожня.pdf").write_bytes(_empty_pdf())
    for path in sorted(HERE.iterdir()):
        if path.suffix != ".py":
            print(f"{path.name}: {path.stat().st_size} bytes")


def _archive_text(path: Path) -> str:
    transcripts = TranscriptStore.load(PROJECT_ROOT / "data" / "transcripts")
    return extract_text(path.read_bytes(), path.name, transcripts).text


def _mushroom_version() -> bytes:
    """60% речень путівника (вступ лишається першим) плюс одне нове речення."""
    _title, _, body = _archive_text(GUIDE).strip().partition("\n")
    sentences = [s for s in re.split(r"(?<=[.!?])\s+", " ".join(body.split())) if s]
    assert sentences[0] == GUIDE_INTRO, sentences[0]
    rest = sentences[1:]
    keep = sorted(random.Random(SEED).sample(range(len(rest)), round(0.6 * len(rest))))
    document = docx.Document()
    document.add_paragraph("Гриби хвойного лісу: нова версія")
    document.add_paragraph(GUIDE_INTRO)
    for i in keep:
        document.add_paragraph(rest[i])
    document.add_paragraph(GUIDE_NEW_SENTENCE)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _scan(font_path: Path) -> bytes:
    """Аркуш A5 при 200 dpi, текст картинкою, без текстового шару."""
    image = Image.new("L", (1169, 1654), color=255)
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype(str(font_path), 42)
    for i, line in enumerate(SCAN_LINES):
        draw.text((120, 200 + i * 90), line, fill=0, font=font)
    buffer = io.BytesIO()
    image.save(buffer, format="PDF", resolution=200)
    return buffer.getvalue()


def _map() -> str:
    text = _archive_text(MAP)
    assert text.count(ON_EDGE) == 1
    return text.replace(ON_EDGE, "На узліссі — хвойний ліс.")


def _too_long() -> str:
    """Рівно 4001 символ без пробілів на краях: ліміт рахується після strip()."""
    sentence = "Ягоди збирайте зранку, поки роса. "
    text = (sentence * (4001 // len(sentence) + 1))[:4001]
    return text[:-1] + "." if text[-1].isspace() else text


def _empty_pdf() -> bytes:
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=420, height=595)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


if __name__ == "__main__":
    main()
